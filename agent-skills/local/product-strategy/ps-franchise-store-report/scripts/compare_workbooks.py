#!/usr/bin/env python3
"""比较两个汇报工作簿的共同汇报页单元格，不导出明细。"""

from __future__ import annotations

import argparse
from pathlib import Path
import json
import math

from openpyxl import load_workbook


DEFAULT_CELLS = (
    "B21,C21,D21,E21,F21,I21,M21,O21,Q21,S21,X21,Y21,"
    "B22,C22,D22,E22,F22,I22,M22,O22,Q22,S22,X22,Y22,"
    "B23,C23,D23,E23,F23,I23,M23,O23,Q23,S23,X23,Y23,"
    "T21,U21,V21,W21,T22,U22,V22,W22,T23,U23,V23,W23"
)


def _number(value: object) -> float | None:
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def compare(reference: Path, candidate: Path, sheet: str, cells: list[str]) -> dict:
    old_ws = load_workbook(reference, data_only=True, read_only=True)[sheet]
    new_ws = load_workbook(candidate, data_only=True, read_only=True)[sheet]
    differences = []
    common_numeric = 0
    for address in cells:
        old = old_ws[address].value
        new = new_ws[address].value
        old_number = _number(old)
        new_number = _number(new)
        if old_number is not None and new_number is not None:
            common_numeric += 1
            delta = new_number - old_number
            relative = None if old_number == 0 else delta / old_number
            if not math.isclose(delta, 0.0, abs_tol=1e-9):
                differences.append({"cell": address, "reference": old_number, "candidate": new_number, "delta": delta, "relative": relative})
        elif old != new:
            differences.append({"cell": address, "reference": old, "candidate": new, "delta": None, "relative": None})
    return {"reference": str(reference), "candidate": str(candidate), "sheet": sheet, "common_numeric": common_numeric, "differences": differences}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--reference", required=True, type=Path)
    parser.add_argument("--candidate", required=True, type=Path)
    parser.add_argument("--sheet", default="结论-组货口径-结果汇报版本")
    parser.add_argument("--cells", default=DEFAULT_CELLS)
    args = parser.parse_args(argv)
    cells = [item for item in args.cells.split(",") if item]
    result = compare(args.reference, args.candidate, args.sheet, cells)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
