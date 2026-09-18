"""Integration tests for O2O Store Benefit Report pipeline (grouping-sets aggregation)."""

from __future__ import annotations

import os
import sys
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace

import openpyxl
import pytest

SKILL_ROOT = Path(__file__).resolve().parent
REPO_ROOT = SKILL_ROOT.parents[4]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))
if str(SKILL_ROOT) not in sys.path:
    sys.path.insert(0, str(SKILL_ROOT))

from helper import (  # noqa: E402
    O2OStoreBenefitRunner,
    _round_significant,
    format_card_summary,
)
from sql_builder import render_dws_and_ads_sql  # noqa: E402


@pytest.fixture
def runner_20260907() -> O2OStoreBenefitRunner:
    """Runner for the 20260907 cutoff using the default H baseline."""
    env_file = os.environ.get("ODPS_ENV_FILE")
    if not env_file:
        pytest.skip("ODPS_ENV_FILE is not set")
    if not Path(env_file).exists():
        pytest.skip(f"ODPS env file not found: {env_file}")
    os.environ["ODPS_ENV_FILE"] = env_file
    return O2OStoreBenefitRunner(
        cutoff_date="20260907",
        version_date="20260705",
        baseline_type="H",
        card_scope="restored_11",
    )


@pytest.fixture
def runner_20260907_raw() -> O2OStoreBenefitRunner:
    """Runner for the default reason_zfl four-card export."""
    env_file = os.environ.get("ODPS_ENV_FILE")
    if not env_file:
        pytest.skip("ODPS_ENV_FILE is not set")
    if not Path(env_file).exists():
        pytest.skip(f"ODPS env file not found: {env_file}")
    os.environ["ODPS_ENV_FILE"] = env_file
    return O2OStoreBenefitRunner(
        cutoff_date="20260907",
        version_date="20260705",
        baseline_type="H",
    )


def test_ads_card_table_produces_11_cards_with_grouping_sets(
    runner_20260907: O2OStoreBenefitRunner,
) -> None:
    """grouping sets must yield 11 cards: 4 所有重点门店 + 4 中心店 + 3 O2O其他重点店."""
    sql = f"""
    select store_group, strategy_tag, count(*) as row_cnt
    from {runner_20260907.target_project}.ads_o2o_key_store_benefit_card_df
    where pt = {runner_20260907.cutoff_date}
      and tag_source = 'restored'
    group by store_group, strategy_tag
    order by
        case when store_group = '所有重点门店' then 1
             when store_group = '中心店' then 2
             when store_group = 'O2O其他重点店' then 3 else 4 end,
        case when strategy_tag = '城市top500品' then 1
             when strategy_tag = '城市top200' then 2
             when strategy_tag = '跨渠道top80' then 3
             when strategy_tag = 'o2o中心店品' then 4 else 5 end
    """
    with runner_20260907.odps.execute_sql(sql).open_reader() as reader:
        cards = [(rec.store_group, rec.strategy_tag, rec.row_cnt) for rec in reader]

    assert len(cards) == 11, f"expected 11 cards, got {len(cards)}"
    assert all(row_cnt == 6 for _, _, row_cnt in cards), (
        "every card must carry 6 metric rows"
    )

    groups_in_order = [g for g, _, _ in cards]
    assert groups_in_order[:4] == ["所有重点门店"] * 4
    assert groups_in_order[4:8] == ["中心店"] * 4
    assert groups_in_order[8:11] == ["O2O其他重点店"] * 3

    tags_all_key_stores = {tag for g, tag, _ in cards if g == "所有重点门店"}
    assert tags_all_key_stores == {
        "城市top500品",
        "城市top200",
        "跨渠道top80",
        "o2o中心店品",
    }

    tags_other_key_stores = {tag for g, tag, _ in cards if g == "O2O其他重点店"}
    assert "o2o中心店品" not in tags_other_key_stores, (
        "honeycomb tag is center-store exclusive"
    )


def test_all_key_stores_equals_sum_of_center_and_other(
    runner_20260907: O2OStoreBenefitRunner,
) -> None:
    """所有重点门店 must equal 中心店 + O2O其他重点店 for sales/margin (additive metrics)."""
    sql = f"""
    select store_group, strategy_tag, metric_name, all_post
    from {runner_20260907.target_project}.ads_o2o_key_store_benefit_card_df
    where pt = {runner_20260907.cutoff_date}
      and tag_source = 'restored'
      and metric_name in ('销售额', '毛利额')
    """
    with runner_20260907.odps.execute_sql(sql).open_reader() as reader:
        rows = list(reader)

    by_key: dict[tuple[str, str], dict[str, Decimal]] = {}
    for row in rows:
        key = (str(row.strategy_tag), str(row.metric_name))
        by_key.setdefault(key, {})[str(row.store_group)] = Decimal(row.all_post)

    assert by_key, "ads card table returned no sales/margin rows"

    for (tag, metric), values in by_key.items():
        assert "所有重点门店" in values
        assert "中心店" in values
        if tag == "o2o中心店品":
            # Honeycomb products are exclusive to center stores; no double count.
            assert "O2O其他重点店" not in values
            assert values["所有重点门店"] == values["中心店"], (
                f"{tag}/{metric}: 所有重点门店 must equal 中心店 for the honeycomb-only tag"
            )
        else:
            assert "O2O其他重点店" in values
            expected = values["中心店"] + values["O2O其他重点店"]
            actual = values["所有重点门店"]
            assert abs(actual - expected) < Decimal("0.1"), (
                f"{tag}/{metric}: 所有重点门店={actual} != 中心店+O2O其他重点店={expected}"
            )


def test_excel_export_renders_11_ordered_card_banners(
    runner_20260907: O2OStoreBenefitRunner,
) -> None:
    """Excel export must render 11 card banners in 所有重点门店→中心店→O2O其他重点店 order."""
    excel_path = Path(runner_20260907.export_excel())
    assert excel_path.exists()

    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    assert ws is not None

    card_headers = [
        cell.value
        for row in ws.iter_rows(min_row=1, max_row=200, min_col=1, max_col=1)
        for cell in row
        if isinstance(cell.value, str)
        and cell.value.startswith("【卡片")
        and cell.value.endswith("效益评估")
    ]

    assert len(card_headers) == 11, f"expected 11 card banners, got {len(card_headers)}"
    assert all("所有重点门店" in h for h in card_headers[0:4])
    assert all("中心店" in h and "所有重点门店" not in h for h in card_headers[4:8])
    assert all("O2O其他重点店" in h for h in card_headers[8:11])


def test_default_reason_zfl_export_renders_4_cards(
    runner_20260907_raw: O2OStoreBenefitRunner,
) -> None:
    """The default raw scope exports four all-key-store cards."""
    sql = f"""
    select store_group, strategy_tag, count(*) as row_cnt
    from {runner_20260907_raw.target_project}.ads_o2o_key_store_benefit_card_df
    where pt = {runner_20260907_raw.cutoff_date}
      and tag_source = 'raw'
      and store_group = '所有重点门店'
    group by store_group, strategy_tag
    order by strategy_tag
    """
    with runner_20260907_raw.odps.execute_sql(sql).open_reader() as reader:
        cards = [(rec.store_group, rec.strategy_tag, rec.row_cnt) for rec in reader]

    assert len(cards) == 4
    assert all(group == "所有重点门店" and row_cnt == 6 for group, _, row_cnt in cards)
    assert {tag for _, tag, _ in cards} == {
        "城市top500品",
        "城市top200",
        "跨渠道top80",
        "o2o中心店品",
    }

    excel_path = Path(runner_20260907_raw.export_excel())
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    assert "基期: 去年同期" in ws["A1"].value
    card_headers = [
        cell.value
        for row in ws.iter_rows(min_row=1, max_row=100, min_col=1, max_col=1)
        for cell in row
        if isinstance(cell.value, str)
        and cell.value.startswith("【卡片")
        and cell.value.endswith("效益评估")
    ]
    assert len(card_headers) == 4
    assert all("所有重点门店" in header for header in card_headers)


def test_all_scope_keeps_tag_sources_as_separate_cards(
    runner_20260907: O2OStoreBenefitRunner,
) -> None:
    """The all scope must not merge raw and restored cards with the same labels."""
    runner_20260907.card_scope = "all"
    sql = f"""
    select tag_source, store_group, strategy_tag, count(*) as row_cnt
    from {runner_20260907.target_project}.ads_o2o_key_store_benefit_card_df
    where pt = {runner_20260907.cutoff_date}
    group by tag_source, store_group, strategy_tag
    """
    with runner_20260907.odps.execute_sql(sql).open_reader() as reader:
        cards = [
            (rec.tag_source, rec.store_group, rec.strategy_tag, rec.row_cnt)
            for rec in reader
        ]

    assert len(cards) == 22
    assert all(row_cnt == 6 for *_, row_cnt in cards)

    excel_path = Path(runner_20260907.export_excel())
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active
    card_headers = [
        cell.value
        for row in ws.iter_rows(min_row=1, max_row=500, min_col=1, max_col=1)
        for cell in row
        if isinstance(cell.value, str)
        and cell.value.startswith("【卡片")
        and cell.value.endswith("效益评估")
    ]
    assert len(card_headers) == 22
    assert any("reason_zfl" in header for header in card_headers)
    assert any("is_3he1_fl" in header for header in card_headers)


def test_all_store_honeycomb_uses_center_catalog_metrics(
    runner_20260907: O2OStoreBenefitRunner,
) -> None:
    """The center-only honeycomb card must not use an all-store catalog numerator."""
    sql = f"""
    select tag_source, period_type, store_group, cata_item_cnt, cata_dx_item_cnt
    from {runner_20260907.target_project}.dws_o2o_key_store_benefit_period_summary_df
    where pt = {runner_20260907.cutoff_date}
      and platform_name = 'all'
      and strategy_tag = 'o2o中心店品'
      and period_type in ('C', 'H')
    """
    with runner_20260907.odps.execute_sql(sql).open_reader() as reader:
        rows = {
            (rec.tag_source, rec.period_type, rec.store_group): (
                rec.cata_item_cnt,
                rec.cata_dx_item_cnt,
            )
            for rec in reader
        }

    for tag_source in ("raw", "restored"):
        for period_type in ("C", "H"):
            assert (
                rows[(tag_source, period_type, "所有重点门店")]
                == rows[(tag_source, period_type, "中心店")]
            )


def test_ads_sql_uses_effective_store_scope_for_share_metrics() -> None:
    _, sql = render_dws_and_ads_sql(
        {"ver_id": 20260705},
        cutoff_date="20260907",
        baseline_type="H",
        target_project="project",
    )

    assert "tmp_o2o_channel_sales_detail_df_2604" in sql
    assert "total_pay_amt" in sql
    assert "total_margin_account" in sql
    assert "d.strategy_tag = 'o2o中心店品'" in sql
    assert "then '中心店'" in sql
    assert "'销售占比' as metric_name" in sql
    assert "'毛利占比' as metric_name" in sql


def test_round_significant_uses_two_sigfigs() -> None:
    assert _round_significant("1.76176714") == "1.8"
    assert _round_significant("0.00981594") == "0.0098"


def test_card_summary_formats_rates_as_percentages() -> None:
    items = [
        {"metric_name": "销售额", "all_diff_ratio": 0.2727},
        {"metric_name": "毛利额", "all_diff_ratio": 0.2},
        {"metric_name": "销售占比", "all_diff": 0.00981594, "all_diff_ratio": None},
        {"metric_name": "毛利占比", "all_diff": 1.76176714, "all_diff_ratio": None},
        {
            "metric_name": "城市top500品动销率",
            "all_diff": -18.56,
            "all_diff_ratio": None,
        },
        {
            "metric_name": "目录内动销率",
            "all_diff": -1.16,
            "all_diff_ratio": None,
        },
    ]

    summary = format_card_summary(
        card_idx=1,
        store_group="所有重点门店",
        strategy_tag="城市top500品",
        source_suffix="",
        items=items,
    )
    assert "标签动销率变化 -18.56%" in summary
    assert "全店目录内动销率变化 -1.16%" in summary
    assert "毛利额变化 +20.00%" in summary
    assert "销售占比变化 +0.0098%" in summary
    assert "毛利占比变化 +1.8%" in summary


def test_honeycomb_share_uses_center_store_denominator(
    runner_20260907: O2OStoreBenefitRunner,
) -> None:
    """The all-store honeycomb card uses the center-store overall denominator."""
    sql = f"""
    select store_group, metric_name, all_post, all_pre
    from {runner_20260907.target_project}.ads_o2o_key_store_benefit_card_df
    where pt = {runner_20260907.cutoff_date}
      and tag_source = 'restored'
      and strategy_tag = 'o2o中心店品'
      and metric_name in ('销售占比', '毛利占比')
      and store_group in ('所有重点门店', '中心店')
    """
    with runner_20260907.odps.execute_sql(sql).open_reader() as reader:
        rows = {
            (rec.metric_name, rec.store_group): (rec.all_post, rec.all_pre)
            for rec in reader
        }

    for metric_name in ("销售占比", "毛利占比"):
        assert rows[(metric_name, "所有重点门店")] == rows[(metric_name, "中心店")]


def test_default_export_queries_raw_cards(tmp_path: Path) -> None:
    columns = [
        "store_group",
        "strategy_tag",
        "order_seq",
        "metric_name",
        "all_post",
        "all_pre",
        "all_diff",
        "all_diff_ratio",
        "o2o_post",
        "o2o_pre",
        "o2o_diff",
        "o2o_diff_ratio",
        "offline_post",
        "offline_pre",
        "offline_diff",
        "offline_diff_ratio",
        "c_avg_tag_items",
        "l_avg_tag_items",
        "c_avg_tag_dx",
        "l_avg_tag_dx",
        "c_avg_cata_items",
        "l_avg_cata_items",
        "c_avg_cata_dx",
        "l_avg_cata_dx",
        "remark",
        "tag_source",
    ]
    values = {column: 1 for column in columns}
    values.update(
        store_group="所有重点门店",
        strategy_tag="城市top500品",
        order_seq=1,
        metric_name="销售占比",
        all_post=1.76176714,
        all_pre=1.012,
        all_diff=0.74976714,
        all_diff_ratio=None,
        remark="",
        tag_source="raw",
    )
    record = [values[column] for column in columns]

    class FakeReader:
        _schema = SimpleNamespace(
            columns=[SimpleNamespace(name=column) for column in columns]
        )

        def open_reader(self) -> "FakeReader":
            return self

        def __enter__(self) -> "FakeReader":
            return self

        def __exit__(self, *_: object) -> None:
            return None

        def __iter__(self):
            return iter([record])

    class FakeODPS:
        sql = ""

        def execute_sql(self, sql: str) -> FakeReader:
            self.sql = sql
            return FakeReader()

    runner = object.__new__(O2OStoreBenefitRunner)
    runner.odps = FakeODPS()
    runner.target_project = "project"
    runner.cutoff_date = 20260907
    runner.card_scope = "raw_all_stores"
    runner.baseline_type = "H"
    runner.export_dir = tmp_path

    runner.export_excel()
    workbook = openpyxl.load_workbook(
        tmp_path / "o2o_key_store_benefit_cards_20260907.xlsx",
        data_only=True,
    )
    sheet = workbook["O2O效益评估卡片"]
    percentage_rows = [
        row
        for row in sheet.iter_rows(min_col=1, max_col=2, values_only=True)
        if row[0] == "销售占比(%)"
    ]
    assert percentage_rows == [("销售占比(%)", 1.8)]
    percentage_row = next(
        row
        for row in sheet.iter_rows(min_col=1, max_col=2)
        if row[0].value == "销售占比(%)"
    )
    assert percentage_row[1].number_format == '0.0"%"'

    assert "tag_source = 'raw'" in runner.odps.sql
    assert "store_group = '所有重点门店'" in runner.odps.sql
