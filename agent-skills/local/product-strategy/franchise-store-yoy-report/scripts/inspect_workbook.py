#!/usr/bin/env python3
"""检查加盟店同比汇报工作簿的结构、说明和图表包。"""

from __future__ import annotations

import argparse
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


def inspect(path: Path) -> list[str]:
    errors: list[str] = []
    workbook = load_workbook(path, data_only=True, read_only=False)
    required = {"结论-组货口径-结果汇报版本", "省区明细"}
    missing = required.difference(workbook.sheetnames)
    if missing:
        errors.append(f"missing sheets: {sorted(missing)}")
    if "结论-组货口径-结果汇报版本" in workbook.sheetnames:
        report = workbook["结论-组货口径-结果汇报版本"]
        if not str(report["A2"].value or "").startswith("结果呈现："):
            errors.append("A2 is missing presentation summary")
        if "桑基图说明：" not in str(report["A61"].value or ""):
            errors.append("A61 is missing Sankey explanation")
        if not str(report["A62"].value or "").startswith("口径："):
            errors.append("A62 is missing scope note")
        if report["T20"].value != "门店数" or report["U20"].value != "销售（同比）":
            errors.append("comparable-market headers are not Chinese")
    if "省区明细" in workbook.sheetnames:
        province = workbook["省区明细"]
        if province["A1"].value != "省区维度数据表现：" or province["B2"].value != "统计时间":
            errors.append("province presentation headers are not aligned")
    with ZipFile(path) as archive:
        names = set(archive.namelist())
        chart_names = {name for name in names if name.startswith("xl/charts/chart") and name.endswith(".xml")}
        sankey_media = {name for name in names if "franchise_store_sankey_graph" in name and name.endswith(".png")}
        if len(chart_names) < 3:
            errors.append(f"expected at least 3 chart XML files, got {len(chart_names)}")
        if not sankey_media:
            errors.append("Sankey PNG is missing from workbook package")
    return errors


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    args = parser.parse_args()
    errors = inspect(args.workbook)
    if errors:
        print("\n".join(errors))
        return 1
    print(f"workbook inspection passed: {args.workbook}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
